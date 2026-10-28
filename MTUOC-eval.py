"""
MTUOC-eval Multi-System Edition (v6.5 - Professional English & Multi-Chart Edition)
Features:
  - Full English UI with concise labels and clean spacing.
  - Increased layout width and height for MT Systems configuration.
  - Segment Details Table ordered: ID, System, [Metrics...], Hypothesis, Reference, Source, Diff.
  - Visualizations with dedicated sub-tabs:
      * Individual Bar Chart tab for every evaluated metric (with Copy & Save).
      * Overall Radar Chart comparison tab (normalized multi-metric spider plot).
      * Individual Radar Chart profile tabs for each evaluated system.
  - Fixed Matplotlib DejaVu Sans 'semibold' font warning.
  - Safe, absolute path loading of 'config.yaml'.
  - Sequential GPU evaluation (PyTorch) with automatic CPU fallback.
  - Full export support: LaTeX (booktabs), Excel (.xlsx), CSV.
"""

import sys
import os
import gc
import csv
import difflib
import glob
import math
import warnings
from typing import List, Dict, Any, Optional

# Suppress deprecation and font warnings cleanly
warnings.filterwarnings("ignore", category=DeprecationWarning)
warnings.filterwarnings("ignore", message=".*findfont:.*")

# ===========================================================================
# Compatibility Shim for legacy packages (Python 3.12+ / setuptools)
# ===========================================================================
try:
    import pkg_resources
except ModuleNotFoundError:
    import types
    import importlib.metadata

    pkg_resources = types.ModuleType("pkg_resources")

    class _DummyDistribution:
        def __init__(self, name):
            self.project_name = str(name)
            try:
                self.version = importlib.metadata.version(str(name))
            except Exception:
                self.version = "0.0.0"

    pkg_resources.get_distribution = lambda dist: _DummyDistribution(dist)
    pkg_resources.require = lambda *args, **kwargs: []
    pkg_resources.DistributionNotFound = Exception
    pkg_resources.VersionConflict = Exception
    sys.modules["pkg_resources"] = pkg_resources
# ===========================================================================

import yaml
import torch

from PyQt6.QtCore import Qt, QThread, pyqtSignal
from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QGridLayout, QGroupBox, QLabel, QLineEdit, QPushButton, QCheckBox,
    QProgressBar, QTableWidget, QTableWidgetItem, QFileDialog, QMessageBox,
    QHeaderView, QTabWidget, QScrollArea, QFrame, QDialog, QPlainTextEdit
)

# Matplotlib integration for PyQt6
import matplotlib
matplotlib.use("QtAgg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg as FigureCanvas
import numpy as np

# Optional openpyxl for Excel export
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# Fast Levenshtein distance with pure Python fallback
try:
    import Levenshtein
    def fast_char_distance(s1: str, s2: str) -> int:
        return Levenshtein.distance(s1, s2)
    def fast_word_distance(w1: list, w2: list) -> int:
        return Levenshtein.distance(w1, w2)
except ImportError:
    def fast_char_distance(s1: str, s2: str) -> int:
        if len(s1) < len(s2):
            return fast_char_distance(s2, s1)
        if len(s2) == 0:
            return len(s1)
        prev = list(range(len(s2) + 1))
        for i, c1 in enumerate(s1):
            curr = [i + 1]
            for j, c2 in enumerate(s2):
                ins = prev[j + 1] + 1
                dels = curr[j] + 1
                subs = prev[j] + (c1 != c2)
                curr.append(min(ins, dels, subs))
            prev = curr
        return prev[-1]

    def fast_word_distance(w1: list, w2: list) -> int:
        if len(w1) < len(w2):
            return fast_word_distance(w2, w1)
        if len(w2) == 0:
            return len(w1)
        prev = list(range(len(w2) + 1))
        for i, t1 in enumerate(w1):
            curr = [i + 1]
            for j, t2 in enumerate(w2):
                ins = prev[j + 1] + 1
                dels = curr[j] + 1
                subs = prev[j] + (t1 != t2)
                curr.append(min(ins, dels, subs))
            prev = curr
        return prev[-1]

# SacreBLEU metrics
try:
    import sacrebleu
    from sacrebleu.metrics import BLEU, CHRF, TER
    HAS_SACREBLEU = True
except ImportError:
    HAS_SACREBLEU = False


# ---------------------------------------------------------------------------
# CUDA Memory Management
# ---------------------------------------------------------------------------

def cleanup_cuda():
    """Forces garbage collection and empties CUDA cache."""
    gc.collect()
    if torch.cuda.is_available():
        torch.cuda.empty_cache()


# ---------------------------------------------------------------------------
# Configuration Loader
# ---------------------------------------------------------------------------

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_FILE_PATH = os.path.join(SCRIPT_DIR, "config.yaml")

DEFAULT_FALLBACK_CONFIG = {
    "Hardware": {
        "device": "auto",
        "batch_size_gpu": 32,
        "batch_size_cpu": 8
    },
    "SacreBLEU": {
        "bleu_tokenize": "13a",
        "bleu_smooth_method": "exp",
        "bleu_lowercase": False,
        "chrf_char_order": 6,
        "chrf_word_order": 2,
        "chrf_whitespace": False,
        "ter_normalized": False,
        "ter_no_punct": False,
        "ter_case_sensitive": False
    },
    "Models": {
        "comet_default": "Unbabel/wmt22-comet-da",
        "comet_qe_default": "Unbabel/wmt20-comet-qe-da",
        "bertscore_default": "bert-base-multilingual-cased",
        "sbert_qe_default": "sentence-transformers/LaBSE"
    }
}

def load_application_config(config_path: str = CONFIG_FILE_PATH) -> dict:
    cfg = {}
    if os.path.exists(config_path):
        try:
            with open(config_path, "r", encoding="utf-8") as f:
                loaded = yaml.safe_load(f)
                if isinstance(loaded, dict):
                    cfg = loaded
        except Exception:
            cfg = {}

    updated_needed = False
    if not cfg:
        cfg = DEFAULT_FALLBACK_CONFIG.copy()
        updated_needed = True
    else:
        for sec, default_vals in DEFAULT_FALLBACK_CONFIG.items():
            if sec not in cfg or not isinstance(cfg[sec], dict):
                cfg[sec] = default_vals.copy()
                updated_needed = True
            else:
                for k, v in default_vals.items():
                    current_val = cfg[sec].get(k)
                    if current_val is None or (isinstance(current_val, str) and not current_val.strip()):
                        cfg[sec][k] = v
                        updated_needed = True

    if updated_needed:
        try:
            with open(config_path, "w", encoding="utf-8") as f:
                yaml.dump(cfg, f, sort_keys=False, default_flow_style=False)
        except Exception:
            pass

    return cfg


# ---------------------------------------------------------------------------
# Auxiliary Functions
# ---------------------------------------------------------------------------

def compute_character_diff(ref: str, hyp: str) -> str:
    matcher = difflib.SequenceMatcher(None, ref, hyp)
    parts = []
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == 'equal':
            parts.append(ref[i1:i2])
        elif tag == 'delete':
            parts.append("{-" + ref[i1:i2] + "-}")
        elif tag == 'insert':
            parts.append("{+" + hyp[j1:j2] + "+}")
        elif tag == 'replace':
            parts.append("{-" + ref[i1:i2] + "-}{+" + hyp[j1:j2] + "+}")
    return "".join(parts)


def compute_segment_ed(ref: str, hyp: str) -> Dict[str, float]:
    c_dist = fast_char_distance(hyp, ref)
    ref_len = max(len(ref), 1)
    c_rate = (c_dist / ref_len) * 100.0

    r_words = ref.strip().split()
    h_words = hyp.strip().split()
    w_dist = fast_word_distance(h_words, r_words)
    w_len = max(len(r_words), 1)
    w_rate = (w_dist / w_len) * 100.0

    return {
        "char_ed": c_dist,
        "char_ed_rate": c_rate,
        "word_ed": w_dist,
        "word_ed_rate": w_rate
    }


def resolve_hf_model_checkpoint(model_name_or_path: str) -> str:
    model_name_or_path = model_name_or_path.strip()
    if not model_name_or_path:
        raise ValueError("Model name or path is not specified.")

    if os.path.isfile(model_name_or_path):
        return model_name_or_path

    repo_id = model_name_or_path
    if "huggingface.co/" in repo_id:
        repo_id = repo_id.split("huggingface.co/")[-1].strip("/")

    try:
        from comet import download_model
        path = download_model(repo_id)
        if path:
            if os.path.isfile(path):
                return path
            ckpts = glob.glob(os.path.join(path, "**", "*.ckpt"), recursive=True)
            if ckpts:
                return ckpts[0]
    except Exception:
        pass

    try:
        from huggingface_hub import snapshot_download
        folder = snapshot_download(repo_id=repo_id)
        ckpts = glob.glob(os.path.join(folder, "**", "*.ckpt"), recursive=True)
        if ckpts:
            return ckpts[0]
        for f in os.listdir(folder):
            if f.endswith(".ckpt"):
                return os.path.join(folder, f)
        return folder
    except Exception as e:
        raise RuntimeError(f"Could not download model '{repo_id}' from Hugging Face: {e}")


def extract_sacrebleu_signature(score_obj, metric_obj) -> str:
    try:
        if hasattr(score_obj, "signature") and score_obj.signature:
            return str(score_obj.signature)
    except Exception:
        pass
    try:
        return str(metric_obj.get_signature())
    except Exception:
        pass
    return "nrefs:1|version:sacrebleu"


# ---------------------------------------------------------------------------
# Background Worker Thread (EvaluationWorker)
# ---------------------------------------------------------------------------

class EvaluationWorker(QThread):
    progress_signal = pyqtSignal(str, int)
    finished_signal = pyqtSignal(dict)
    error_signal = pyqtSignal(str)

    def __init__(
        self,
        mode: str,
        src_path: str,
        ref_path: str,
        systems: List[Dict[str, str]],
        selected_metrics: Dict[str, bool],
        config: dict,
        comet_model: str,
        comet_qe_model: str,
        bertscore_model: str,
        sbert_qe_model: str
    ):
        super().__init__()
        self.mode = mode
        self.src_path = src_path
        self.ref_path = ref_path
        self.systems = systems
        self.selected_metrics = selected_metrics
        self.config = config
        self.comet_model = comet_model
        self.comet_qe_model = comet_qe_model
        self.bertscore_model = bertscore_model
        self.sbert_qe_model = sbert_qe_model

    def run(self):
        try:
            self.progress_signal.emit("Initializing environment...", 5)

            hw_cfg = self.config.get("Hardware", {})
            hw_setting = hw_cfg.get("device", "auto").lower()
            if hw_setting == "cuda":
                use_cuda = torch.cuda.is_available()
            elif hw_setting == "cpu":
                use_cuda = False
            else:
                use_cuda = torch.cuda.is_available()

            num_gpus = 1 if use_cuda else 0
            device_str = "cuda" if use_cuda else "cpu"
            batch_size_gpu = hw_cfg.get("batch_size_gpu", 32)
            batch_size_cpu = hw_cfg.get("batch_size_cpu", 8)
            batch_size = batch_size_gpu if use_cuda else batch_size_cpu

            sources = []
            references = []

            if self.mode == "reference":
                if not self.ref_path or not os.path.exists(self.ref_path):
                    raise ValueError("Reference file is required for reference-based evaluation.")
                with open(self.ref_path, "r", encoding="utf-8") as f:
                    references = [line.rstrip("\r\n") for line in f]
                total_lines = len(references)
                if total_lines == 0:
                    raise ValueError("Reference file is empty.")

                if self.src_path and os.path.exists(self.src_path):
                    with open(self.src_path, "r", encoding="utf-8") as f:
                        sources = [line.rstrip("\r\n") for line in f]
                    if len(sources) != total_lines:
                        raise ValueError(f"Line mismatch: reference has {total_lines} lines, source has {len(sources)} lines.")
            else:
                if not self.src_path or not os.path.exists(self.src_path):
                    raise ValueError("Source file is required for referenceless (QE) evaluation.")
                with open(self.src_path, "r", encoding="utf-8") as f:
                    sources = [line.rstrip("\r\n") for line in f]
                total_lines = len(sources)
                if total_lines == 0:
                    raise ValueError("Source file is empty.")

            loaded_systems = []
            for s_info in self.systems:
                name = s_info["name"]
                path = s_info["path"]
                with open(path, "r", encoding="utf-8") as f:
                    hyps = [line.rstrip("\r\n") for line in f]
                if len(hyps) != total_lines:
                    raise ValueError(f"Line mismatch: expected {total_lines} lines, but '{name}' has {len(hyps)} lines.")
                loaded_systems.append({"name": name, "hyps": hyps})

            corpus_summary = {}
            signatures = {}

            bleu_seg_scores = {}
            chrf_seg_scores = {}
            ter_seg_scores = {}
            ed_seg_scores = {}
            comet_seg_scores = {}
            bert_seg_scores = {}
            comet_qe_seg_scores = {}
            sbert_qe_seg_scores = {}

            # ===============================================================
            # PHASE 1: Rule-based Metrics (Reference Mode)
            # ===============================================================
            if self.mode == "reference":
                sb_cfg = self.config.get("SacreBLEU", {})
                need_sacrebleu = any(self.selected_metrics.get(k) for k in ["bleu", "chrf", "ter"])
                if need_sacrebleu and not HAS_SACREBLEU:
                    raise ImportError("sacrebleu is not installed. Run: pip install sacrebleu")

                bleu_obj = BLEU(
                    tokenize=sb_cfg.get("bleu_tokenize", "13a"),
                    smooth_method=sb_cfg.get("bleu_smooth_method", "exp"),
                    lowercase=sb_cfg.get("bleu_lowercase", False),
                    effective_order=True
                ) if self.selected_metrics.get("bleu") else None

                chrf_obj = CHRF(
                    char_order=sb_cfg.get("chrf_char_order", 6),
                    word_order=sb_cfg.get("chrf_word_order", 2),
                    whitespace=sb_cfg.get("chrf_whitespace", False)
                ) if self.selected_metrics.get("chrf") else None

                ter_obj = TER(
                    normalized=sb_cfg.get("ter_normalized", False),
                    no_punct=sb_cfg.get("ter_no_punct", False),
                    case_sensitive=sb_cfg.get("ter_case_sensitive", False)
                ) if self.selected_metrics.get("ter") else None

                self.progress_signal.emit("Evaluating standard metrics (BLEU, chrF++, TER, Edit Distance)...", 10)

                for sys_item in loaded_systems:
                    s_name = sys_item["name"]
                    s_hyps = sys_item["hyps"]

                    if bleu_obj:
                        res = bleu_obj.corpus_score(s_hyps, [references])
                        corpus_summary.setdefault("BLEU", {})[s_name] = round(res.score, 2)
                        if "BLEU" not in signatures:
                            signatures["BLEU"] = extract_sacrebleu_signature(res, bleu_obj)
                        bleu_seg_scores[s_name] = [
                            round(bleu_obj.sentence_score(h, [r]).score, 2)
                            for h, r in zip(s_hyps, references)
                        ]

                    if chrf_obj:
                        res = chrf_obj.corpus_score(s_hyps, [references])
                        corpus_summary.setdefault("chrF++", {})[s_name] = round(res.score, 2)
                        if "chrF++" not in signatures:
                            signatures["chrF++"] = extract_sacrebleu_signature(res, chrf_obj)
                        chrf_seg_scores[s_name] = [
                            round(chrf_obj.sentence_score(h, [r]).score, 2)
                            for h, r in zip(s_hyps, references)
                        ]

                    if ter_obj:
                        res = ter_obj.corpus_score(s_hyps, [references])
                        corpus_summary.setdefault("TER", {})[s_name] = round(res.score, 2)
                        if "TER" not in signatures:
                            signatures["TER"] = extract_sacrebleu_signature(res, ter_obj)
                        ter_seg_scores[s_name] = [
                            round(ter_obj.sentence_score(h, [r]).score, 2)
                            for h, r in zip(s_hyps, references)
                        ]

                    if self.selected_metrics.get("ed_dist"):
                        tot_c_dist = sum(fast_char_distance(h, r) for h, r in zip(s_hyps, references))
                        tot_c_len = sum(max(len(r), 1) for r in references)
                        tot_w_dist = sum(fast_word_distance(h.split(), r.split()) for h, r in zip(s_hyps, references))
                        tot_w_len = sum(max(len(r.split()), 1) for r in references)

                        corpus_summary.setdefault("%EdDist (Char)", {})[s_name] = round((tot_c_dist / tot_c_len) * 100.0, 2)
                        corpus_summary.setdefault("WER (Word)", {})[s_name] = round((tot_w_dist / tot_w_len) * 100.0, 2)

                        signatures["%EdDist (Char)"] = "Levenshtein Character Edit Distance normalized by Reference length (%)"
                        signatures["WER (Word)"] = "Levenshtein Word Error Rate normalized by Reference word count (%)"

                        ed_seg_scores[s_name] = [compute_segment_ed(r, h) for r, h in zip(references, s_hyps)]

            # ===============================================================
            # PHASE 2: COMET (Reference Mode)
            # ===============================================================
            if self.mode == "reference" and self.selected_metrics.get("comet"):
                if not sources:
                    raise ValueError("COMET evaluation requires a source file.")
                if not self.comet_model:
                    raise ValueError("No COMET model specified in config.yaml or input.")

                self.progress_signal.emit(f"Loading COMET ({self.comet_model})...", 30)
                from comet import load_from_checkpoint
                ckpt = resolve_hf_model_checkpoint(self.comet_model)
                comet_inst = load_from_checkpoint(ckpt)

                actual_comet_dev = device_str
                try:
                    for idx, sys_item in enumerate(loaded_systems):
                        s_name = sys_item["name"]
                        pct = 30 + int(((idx + 1) / len(loaded_systems)) * 30)
                        self.progress_signal.emit(f"Evaluating COMET for {s_name} ({actual_comet_dev.upper()})...", pct)

                        data = [{"src": s, "mt": h, "ref": r} for s, h, r in zip(sources, sys_item["hyps"], references)]
                        out = comet_inst.predict(data, batch_size=batch_size, gpus=num_gpus)
                        corpus_summary.setdefault("COMET", {})[s_name] = round(out.system_score, 4)
                        comet_seg_scores[s_name] = [round(sc, 4) for sc in out.scores]
                except Exception as cuda_ex:
                    if use_cuda:
                        self.progress_signal.emit(f"COMET GPU issue ({cuda_ex}). Falling back to CPU...", 35)
                        cleanup_cuda()
                        actual_comet_dev = "cpu"
                        for idx, sys_item in enumerate(loaded_systems):
                            s_name = sys_item["name"]
                            data = [{"src": s, "mt": h, "ref": r} for s, h, r in zip(sources, sys_item["hyps"], references)]
                            out = comet_inst.predict(data, batch_size=batch_size_cpu, gpus=0)
                            corpus_summary.setdefault("COMET", {})[s_name] = round(out.system_score, 4)
                            comet_seg_scores[s_name] = [round(sc, 4) for sc in out.scores]
                    else:
                        raise cuda_ex

                signatures["COMET"] = f"model:{self.comet_model}|device:{actual_comet_dev}"
                del comet_inst
                cleanup_cuda()

            # ===============================================================
            # PHASE 3: BERTScore (Reference Mode)
            # ===============================================================
            if self.mode == "reference" and self.selected_metrics.get("bertscore"):
                if not self.bertscore_model:
                    raise ValueError("No BERTScore model specified in config.yaml or input.")

                self.progress_signal.emit(f"Loading BERTScore ({self.bertscore_model})...", 65)
                import bert_score
                bs_model = self.bertscore_model.strip()
                if "huggingface.co/" in bs_model:
                    bs_model = bs_model.split("huggingface.co/")[-1].strip("/")

                actual_bert_dev = device_str
                try:
                    for idx, sys_item in enumerate(loaded_systems):
                        s_name = sys_item["name"]
                        pct = 65 + int(((idx + 1) / len(loaded_systems)) * 25)
                        self.progress_signal.emit(f"Evaluating BERTScore for {s_name} ({actual_bert_dev.upper()})...", pct)

                        _, _, F1 = bert_score.score(
                            sys_item["hyps"], references, model_type=bs_model,
                            device=actual_bert_dev, batch_size=batch_size, verbose=False
                        )
                        f1_vals = F1.tolist()
                        corpus_summary.setdefault("BERTScore (F1)", {})[s_name] = round(float(sum(f1_vals) / len(f1_vals)), 4)
                        bert_seg_scores[s_name] = [round(val, 4) for val in f1_vals]
                except Exception as cuda_ex:
                    if use_cuda:
                        self.progress_signal.emit(f"BERTScore GPU issue ({cuda_ex}). Falling back to CPU...", 70)
                        cleanup_cuda()
                        actual_bert_dev = "cpu"
                        for idx, sys_item in enumerate(loaded_systems):
                            s_name = sys_item["name"]
                            _, _, F1 = bert_score.score(
                                sys_item["hyps"], references, model_type=bs_model,
                                device="cpu", batch_size=batch_size_cpu, verbose=False
                            )
                            f1_vals = F1.tolist()
                            corpus_summary.setdefault("BERTScore (F1)", {})[s_name] = round(float(sum(f1_vals) / len(f1_vals)), 4)
                            bert_seg_scores[s_name] = [round(val, 4) for val in f1_vals]
                    else:
                        raise cuda_ex

                signatures["BERTScore"] = f"model:{self.bertscore_model}|device:{actual_bert_dev}"
                cleanup_cuda()

            # ===============================================================
            # PHASE 4: COMET-QE (Referenceless Mode)
            # ===============================================================
            if self.mode == "referenceless" and self.selected_metrics.get("comet_qe"):
                if not self.comet_qe_model:
                    raise ValueError("No COMET-QE model specified in config.yaml or input.")

                self.progress_signal.emit(f"Loading COMET-QE ({self.comet_qe_model})...", 20)
                from comet import load_from_checkpoint
                ckpt = resolve_hf_model_checkpoint(self.comet_qe_model)
                comet_qe_inst = load_from_checkpoint(ckpt)

                actual_qe_dev = device_str
                try:
                    for idx, sys_item in enumerate(loaded_systems):
                        s_name = sys_item["name"]
                        pct = 20 + int(((idx + 1) / len(loaded_systems)) * 35)
                        self.progress_signal.emit(f"Evaluating COMET-QE for {s_name} ({actual_qe_dev.upper()})...", pct)

                        data = [{"src": s, "mt": h} for s, h in zip(sources, sys_item["hyps"])]
                        out = comet_qe_inst.predict(data, batch_size=batch_size, gpus=num_gpus)
                        corpus_summary.setdefault("COMET-QE", {})[s_name] = round(out.system_score, 4)
                        comet_qe_seg_scores[s_name] = [round(sc, 4) for sc in out.scores]
                except Exception as cuda_ex:
                    if use_cuda:
                        self.progress_signal.emit(f"COMET-QE GPU issue ({cuda_ex}). Falling back to CPU...", 30)
                        cleanup_cuda()
                        actual_qe_dev = "cpu"
                        for idx, sys_item in enumerate(loaded_systems):
                            s_name = sys_item["name"]
                            data = [{"src": s, "mt": h} for s, h in zip(sources, sys_item["hyps"])]
                            out = comet_qe_inst.predict(data, batch_size=batch_size_cpu, gpus=0)
                            corpus_summary.setdefault("COMET-QE", {})[s_name] = round(out.system_score, 4)
                            comet_qe_seg_scores[s_name] = [round(sc, 4) for sc in out.scores]
                    else:
                        raise cuda_ex

                signatures["COMET-QE"] = f"model:{self.comet_qe_model}|device:{actual_qe_dev}"
                del comet_qe_inst
                cleanup_cuda()

            # ===============================================================
            # PHASE 5: SBERT Cosine Similarity (Referenceless Mode)
            # ===============================================================
            if self.mode == "referenceless" and self.selected_metrics.get("sbert_qe"):
                if not self.sbert_qe_model:
                    raise ValueError("No SBERT model specified in config.yaml or input.")

                self.progress_signal.emit(f"Loading SBERT ({self.sbert_qe_model})...", 60)
                try:
                    from sentence_transformers import SentenceTransformer, util
                except ImportError:
                    raise ImportError("sentence-transformers is not installed. Run: pip install sentence-transformers")

                sbert_model = SentenceTransformer(self.sbert_qe_model, device=device_str)
                signatures["SBERT-QE"] = f"model:{self.sbert_qe_model}|metric:cosine_similarity|device:{device_str}"

                src_embeddings = sbert_model.encode(
                    sources, batch_size=batch_size, convert_to_tensor=True, show_progress_bar=False
                )

                for idx, sys_item in enumerate(loaded_systems):
                    s_name = sys_item["name"]
                    pct = 60 + int(((idx + 1) / len(loaded_systems)) * 30)
                    self.progress_signal.emit(f"Evaluating SBERT Cosine for {s_name}...", pct)

                    hyp_embeddings = sbert_model.encode(
                        sys_item["hyps"], batch_size=batch_size, convert_to_tensor=True, show_progress_bar=False
                    )
                    cos_sims = util.cos_sim(src_embeddings, hyp_embeddings).diagonal()
                    scores_list = [round(float(s), 4) for s in cos_sims]

                    corpus_summary.setdefault("SBERT-QE (Cosine)", {})[s_name] = round(float(sum(scores_list) / len(scores_list)), 4)
                    sbert_qe_seg_scores[s_name] = scores_list

                del sbert_model
                cleanup_cuda()

            # ===============================================================
            # PHASE 6: Assemble Segment Records
            # ===============================================================
            self.progress_signal.emit("Building detailed segment records...", 95)
            all_segment_records = []

            for sys_item in loaded_systems:
                s_name = sys_item["name"]
                s_hyps = sys_item["hyps"]

                for i in range(total_lines):
                    rec = {"seg_id": i + 1, "system": s_name, "metrics": {}}

                    if self.mode == "reference":
                        if s_name in bleu_seg_scores:
                            rec["metrics"]["BLEU"] = bleu_seg_scores[s_name][i]
                        if s_name in chrf_seg_scores:
                            rec["metrics"]["chrF++"] = chrf_seg_scores[s_name][i]
                        if s_name in ter_seg_scores:
                            rec["metrics"]["TER"] = ter_seg_scores[s_name][i]
                        if s_name in ed_seg_scores:
                            ed_i = ed_seg_scores[s_name][i]
                            rec["metrics"]["%EdDist"] = round(ed_i["char_ed_rate"], 2)
                            rec["metrics"]["WER"] = round(ed_i["word_ed_rate"], 2)
                        if s_name in comet_seg_scores:
                            rec["metrics"]["COMET"] = comet_seg_scores[s_name][i]
                        if s_name in bert_seg_scores:
                            rec["metrics"]["BERTScore"] = bert_seg_scores[s_name][i]

                        rec["hypothesis"] = s_hyps[i]
                        rec["reference"] = references[i]
                        rec["source"] = sources[i] if sources else ""
                        rec["diff"] = compute_character_diff(references[i], s_hyps[i])
                    else:
                        if s_name in comet_qe_seg_scores:
                            rec["metrics"]["COMET-QE"] = comet_qe_seg_scores[s_name][i]
                        if s_name in sbert_qe_seg_scores:
                            rec["metrics"]["SBERT-QE"] = sbert_qe_seg_scores[s_name][i]

                        rec["hypothesis"] = s_hyps[i]
                        rec["reference"] = ""
                        rec["source"] = sources[i]
                        rec["diff"] = ""

                    all_segment_records.append(rec)

            self.progress_signal.emit("Evaluation complete.", 100)
            self.finished_signal.emit({
                "mode": self.mode,
                "corpus": corpus_summary,
                "segments": all_segment_records,
                "systems": [s["name"] for s in loaded_systems],
                "signatures": signatures,
                "has_source": bool(sources),
                "has_ref": bool(references)
            })

        except Exception as e:
            cleanup_cuda()
            self.error_signal.emit(str(e))


# ---------------------------------------------------------------------------
# Dynamic System Row Widget (Increased Width for Clean Layout)
# ---------------------------------------------------------------------------

class SystemRowWidget(QWidget):
    remove_requested = pyqtSignal(object)

    def __init__(self, system_index: int, default_name: str = ""):
        super().__init__()
        layout = QHBoxLayout(self)
        layout.setContentsMargins(0, 3, 0, 3)
        layout.setSpacing(8)

        self.lbl = QLabel(f"System #{system_index}:")
        self.lbl.setFixedWidth(80)

        self.name_edit = QLineEdit(default_name or f"System_{system_index}")
        self.name_edit.setPlaceholderText("System Name")
        self.name_edit.setFixedWidth(160)

        self.path_edit = QLineEdit()
        self.path_edit.setPlaceholderText("Path to translation file (*)")

        self.btn_browse = QPushButton("Browse...")
        self.btn_browse.setFixedWidth(85)
        self.btn_browse.clicked.connect(self.browse_file)

        self.btn_remove = QPushButton("Remove")
        self.btn_remove.setFixedWidth(80)
        self.btn_remove.setStyleSheet("color: #b30000;")
        self.btn_remove.clicked.connect(lambda: self.remove_requested.emit(self))

        layout.addWidget(self.lbl)
        layout.addWidget(self.name_edit)
        layout.addWidget(self.path_edit)
        layout.addWidget(self.btn_browse)
        layout.addWidget(self.btn_remove)

    def browse_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "Select Translation File", "", "All Files (*)")
        if path:
            self.path_edit.setText(path)
            if not self.name_edit.text() or self.name_edit.text().startswith("System_"):
                base_name = os.path.splitext(os.path.basename(path))[0]
                self.name_edit.setText(base_name)

    def get_data(self) -> Dict[str, str]:
        return {
            "name": self.name_edit.text().strip() or "Unnamed_System",
            "path": self.path_edit.text().strip()
        }


# ---------------------------------------------------------------------------
# Single Metric Bar Plot Page
# ---------------------------------------------------------------------------

class SingleMetricPlotPage(QWidget):
    def __init__(self, metric_name: str, systems: List[str], scores: List[float]):
        super().__init__()
        self.metric_name = metric_name

        layout = QVBoxLayout(self)
        layout.setContentsMargins(10, 10, 10, 10)

        bar = QHBoxLayout()
        title_lbl = QLabel(f"<b>Metric: {metric_name}</b>")
        title_lbl.setStyleSheet("font-size: 13px;")

        btn_copy = QPushButton("Copy to Clipboard")
        btn_copy.clicked.connect(self.copy_to_clipboard)

        btn_save = QPushButton("Save as PNG...")
        btn_save.clicked.connect(self.save_as_png)

        bar.addWidget(title_lbl)
        bar.addStretch()
        bar.addWidget(btn_copy)
        bar.addWidget(btn_save)
        layout.addLayout(bar)

        self.fig, self.ax = plt.subplots(figsize=(8.5, 4.8))
        self.canvas = FigureCanvas(self.fig)
        layout.addWidget(self.canvas)

        palette = plt.cm.tab10(np.linspace(0, 1, max(len(systems), 10)))
        bars = self.ax.bar(
            systems, scores,
            color=[palette[i % len(palette)] for i in range(len(systems))],
            width=0.5, edgecolor='black', linewidth=0.8
        )
        self.ax.set_title(f"{metric_name} Comparison Across Systems", fontsize=12, fontweight="bold", pad=12)
        self.ax.set_ylabel("Score", fontsize=10)
        self.ax.grid(axis='y', linestyle='--', alpha=0.5)

        for bar_item in bars:
            h = bar_item.get_height()
            if isinstance(h, (int, float)):
                fmt = f"{h:.2f}" if abs(h) >= 1 else f"{h:.4f}"
                self.ax.annotate(fmt,
                                 xy=(bar_item.get_x() + bar_item.get_width() / 2, h),
                                 xytext=(0, 4),
                                 textcoords="offset points",
                                 ha='center', va='bottom', fontsize=9, fontweight='bold')

        y_vals = [v for v in scores if isinstance(v, (int, float))]
        if y_vals:
            min_y = min(y_vals)
            max_y = max(y_vals)
            span = max_y - min_y if max_y != min_y else abs(max_y) * 0.2
            margin = span * 0.2 if span != 0 else 1.0
            self.ax.set_ylim([max(0, min_y - margin/2), max_y + margin])

        self.fig.tight_layout()
        self.canvas.draw()

    def copy_to_clipboard(self):
        pixmap = self.canvas.grab()
        QApplication.clipboard().setPixmap(pixmap)
        QMessageBox.information(self, "Copied", f"'{self.metric_name}' chart copied to clipboard!")

    def save_as_png(self):
        clean_name = self.metric_name.replace(" ", "_").replace("%", "pct").replace("/", "_")
        path, _ = QFileDialog.getSaveFileName(self, f"Save {self.metric_name} Chart", f"{clean_name}_bar_chart.png", "PNG Image (*.png);;All Files (*)")
        if path:
            self.fig.savefig(path, dpi=300, bbox_inches="tight")
            QMessageBox.information(self, "Saved", f"Chart saved to:\n{path}")


# ---------------------------------------------------------------------------
# Radar / Spider Plot Page
# ---------------------------------------------------------------------------

class RadarPlotPage(QWidget):
    def __init__(self, title: str, metrics: List[str], series_dict: Dict[str, List[float]]):
        """
        title: Title of the radar plot
        metrics: list of metric names (axes)
        series_dict: {series_name: [val1, val2, ...]} normalized or formatted
        """
        super().__init__()
        self.chart_title = title

        layout = QVBoxLayout(self)
        layout.setContentsMargins(10, 10, 10, 10)

        bar = QHBoxLayout()
        title_lbl = QLabel(f"<b>Radar Chart: {title}</b>")
        title_lbl.setStyleSheet("font-size: 13px;")

        btn_copy = QPushButton("Copy to Clipboard")
        btn_copy.clicked.connect(self.copy_to_clipboard)

        btn_save = QPushButton("Save as PNG...")
        btn_save.clicked.connect(self.save_as_png)

        bar.addWidget(title_lbl)
        bar.addStretch()
        bar.addWidget(btn_copy)
        bar.addWidget(btn_save)
        layout.addLayout(bar)

        self.fig = plt.figure(figsize=(8.5, 5.0))
        self.ax = self.fig.add_subplot(111, polar=True)
        self.canvas = FigureCanvas(self.fig)
        layout.addWidget(self.canvas)

        num_vars = len(metrics)
        angles = np.linspace(0, 2 * np.pi, num_vars, endpoint=False).tolist()
        angles += angles[:1]

        palette = plt.cm.tab10(np.linspace(0, 1, max(len(series_dict), 10)))

        for idx, (s_name, values) in enumerate(series_dict.items()):
            vals = values + values[:1]
            color = palette[idx % len(palette)]
            self.ax.plot(angles, vals, linewidth=2, linestyle='solid', label=s_name, color=color)
            self.ax.fill(angles, vals, color=color, alpha=0.15)

        self.ax.set_xticks(angles[:-1])
        self.ax.set_xticklabels(metrics, fontsize=9, fontweight='bold')
        self.ax.set_title(title, fontsize=12, fontweight='bold', pad=15)
        self.ax.legend(loc='upper right', bbox_to_anchor=(1.3, 1.1), fontsize=9)
        self.ax.grid(True, linestyle='--', alpha=0.6)

        self.fig.tight_layout()
        self.canvas.draw()

    def copy_to_clipboard(self):
        pixmap = self.canvas.grab()
        QApplication.clipboard().setPixmap(pixmap)
        QMessageBox.information(self, "Copied", f"'{self.chart_title}' radar chart copied to clipboard!")

    def save_as_png(self):
        clean_name = self.chart_title.replace(" ", "_").replace("%", "pct").replace("/", "_")
        path, _ = QFileDialog.getSaveFileName(self, "Save Radar Chart", f"{clean_name}_radar.png", "PNG Image (*.png);;All Files (*)")
        if path:
            self.fig.savefig(path, dpi=300, bbox_inches="tight")
            QMessageBox.information(self, "Saved", f"Radar chart saved to:\n{path}")


# ---------------------------------------------------------------------------
# Visualizations Widget (Tabs for Bar Charts & Radar Charts)
# ---------------------------------------------------------------------------

class VisualizationWidget(QWidget):
    def __init__(self):
        super().__init__()
        self.layout = QVBoxLayout(self)

        self.lbl_notice = QLabel("Comparative charts require at least 2 systems to be evaluated.")
        self.lbl_notice.setStyleSheet("color: #666; font-style: italic; font-size: 13px; margin: 30px;")
        self.lbl_notice.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.layout.addWidget(self.lbl_notice)

        self.metric_tabs = QTabWidget()
        self.layout.addWidget(self.metric_tabs)
        self.metric_tabs.hide()

    def update_plots(self, corpus_data: Dict[str, Dict[str, Any]], systems: List[str]):
        self.metric_tabs.clear()

        if len(systems) <= 1:
            self.lbl_notice.setText("Only 1 system evaluated. Add 2 or more systems to compare metrics.")
            self.lbl_notice.show()
            self.metric_tabs.hide()
            return

        numeric_metrics = [m for m, s_dict in corpus_data.items() if any(isinstance(v, (int, float)) for v in s_dict.values())]
        if not numeric_metrics:
            self.lbl_notice.setText("No numeric metrics available to plot.")
            self.lbl_notice.show()
            self.metric_tabs.hide()
            return

        self.lbl_notice.hide()
        self.metric_tabs.show()

        # 1. Bar Chart Tabs: One tab per metric
        for metric in numeric_metrics:
            scores = [corpus_data[metric].get(s, 0.0) for s in systems]
            page = SingleMetricPlotPage(metric, systems, scores)
            self.metric_tabs.addTab(page, f"Bar: {metric}")

        # 2. Overall Radar Chart (Normalized 0-100% across metrics)
        # Normalization allows comparing BLEU (0-100) and COMET (0-1) on the same spider web
        norm_series = {s: [] for s in systems}
        for metric in numeric_metrics:
            raw_scores = [corpus_data[metric].get(s, 0.0) for s in systems]
            min_v = min(raw_scores)
            max_v = max(raw_scores)
            diff = max_v - min_v
            for s in systems:
                val = corpus_data[metric].get(s, 0.0)
                if diff > 0:
                    norm_score = ((val - min_v) / diff) * 100.0
                else:
                    norm_score = 100.0
                norm_series[s].append(round(norm_score, 1))

        radar_all_page = RadarPlotPage(
            title="All Systems Relative Profile (Normalized 0-100%)",
            metrics=numeric_metrics,
            series_dict=norm_series
        )
        self.metric_tabs.addTab(radar_all_page, "Radar: All Systems")

        # 3. Individual Radar Chart per system (showing raw metric scores)
        for s in systems:
            sys_raw = [corpus_data[m].get(s, 0.0) for m in numeric_metrics]
            radar_sys_page = RadarPlotPage(
                title=f"{s} Metric Footprint",
                metrics=numeric_metrics,
                series_dict={s: sys_raw}
            )
            self.metric_tabs.addTab(radar_sys_page, f"Radar: {s}")


# ---------------------------------------------------------------------------
# LaTeX Table Export Dialog
# ---------------------------------------------------------------------------

class LatexExportDialog(QDialog):
    def __init__(self, latex_code: str, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Export to LaTeX (booktabs format)")
        self.resize(750, 500)
        self.latex_code = latex_code

        layout = QVBoxLayout(self)

        info_lbl = QLabel(
            "<b>LaTeX Table Code (add <code>\\usepackage{booktabs}</code> to your preamble):</b>"
        )
        layout.addWidget(info_lbl)

        self.txt_edit = QPlainTextEdit(self.latex_code)
        self.txt_edit.setStyleSheet("font-family: Consolas, 'Courier New', monospace; font-size: 11px;")
        layout.addWidget(self.txt_edit)

        h_btns = QHBoxLayout()
        btn_copy = QPushButton("Copy to Clipboard")
        btn_copy.setStyleSheet("font-weight: bold; background-color: #0066cc; color: white; padding: 6px;")
        btn_copy.clicked.connect(self.copy_clipboard)
        h_btns.addWidget(btn_copy)

        btn_save = QPushButton("Save as .tex...")
        btn_save.clicked.connect(self.save_tex)
        h_btns.addWidget(btn_save)

        btn_close = QPushButton("Close")
        btn_close.clicked.connect(self.accept)
        h_btns.addWidget(btn_close)

        layout.addLayout(h_btns)

    def copy_clipboard(self):
        QApplication.clipboard().setText(self.txt_edit.toPlainText())
        QMessageBox.information(self, "Copied", "LaTeX table copied to clipboard!")

    def save_tex(self):
        path, _ = QFileDialog.getSaveFileName(self, "Save LaTeX Table", "mt_evaluation_table.tex", "TeX Files (*.tex);;All Files (*)")
        if path:
            with open(path, "w", encoding="utf-8") as f:
                f.write(self.txt_edit.toPlainText())
            QMessageBox.information(self, "Saved", f"File saved successfully to:\n{path}")


# ---------------------------------------------------------------------------
# Main Application Window
# ---------------------------------------------------------------------------

class MTUOCEvalApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.config = load_application_config(CONFIG_FILE_PATH)

        self.setWindowTitle("MTUOC-eval v6.5 — Machine Translation Evaluation Platform")
        self.resize(1220, 900)
        self.results_data: Optional[Dict[str, Any]] = None

        self.ref_system_rows: List[SystemRowWidget] = []
        self.qe_system_rows: List[SystemRowWidget] = []

        self.init_ui()

    def init_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        main_layout = QVBoxLayout(central)

        self.tabs = QTabWidget()
        main_layout.addWidget(self.tabs)

        models_cfg = self.config.get("Models", {})

        # -------------------------------------------------------------------
        # TAB 1: Reference-based Evaluation
        # -------------------------------------------------------------------
        tab_ref = QWidget()
        layout_ref = QVBoxLayout(tab_ref)

        box_ref_files = QGroupBox("Input Files")
        grid_ref_files = QGridLayout(box_ref_files)

        grid_ref_files.addWidget(QLabel("Source file (Optional, needed for COMET):"), 0, 0)
        self.txt_ref_src = QLineEdit()
        self.txt_ref_src.setPlaceholderText("Select optional source text file (*)")
        grid_ref_files.addWidget(self.txt_ref_src, 0, 1)
        btn_browse_src = QPushButton("Browse...")
        btn_browse_src.clicked.connect(lambda: self.browse_file(self.txt_ref_src, "Select Source File"))
        grid_ref_files.addWidget(btn_browse_src, 0, 2)

        lbl_ref = QLabel("<b>Reference file (* Required):</b>")
        grid_ref_files.addWidget(lbl_ref, 1, 0)
        self.txt_ref_ref = QLineEdit()
        self.txt_ref_ref.setPlaceholderText("Select human reference translation file (*)")
        grid_ref_files.addWidget(self.txt_ref_ref, 1, 1)
        btn_browse_ref = QPushButton("Browse...")
        btn_browse_ref.clicked.connect(lambda: self.browse_file(self.txt_ref_ref, "Select Reference File"))
        grid_ref_files.addWidget(btn_browse_ref, 1, 2)

        layout_ref.addWidget(box_ref_files)

        # MT Systems box with generous width and height
        box_ref_sys = QGroupBox("MT Systems to Evaluate")
        vbox_ref_sys = QVBoxLayout(box_ref_sys)
        scroll_ref = QScrollArea()
        scroll_ref.setWidgetResizable(True)
        scroll_ref.setMinimumHeight(150)
        scroll_ref.setFrameShape(QFrame.Shape.NoFrame)
        self.cont_ref_sys = QWidget()
        self.layout_ref_sys = QVBoxLayout(self.cont_ref_sys)
        self.layout_ref_sys.setAlignment(Qt.AlignmentFlag.AlignTop)
        scroll_ref.setWidget(self.cont_ref_sys)
        vbox_ref_sys.addWidget(scroll_ref)

        btn_add_ref_sys = QPushButton("+ Add Another System")
        btn_add_ref_sys.setStyleSheet("font-weight: bold; padding: 6px;")
        btn_add_ref_sys.clicked.connect(lambda: self.add_system_row(self.layout_ref_sys, self.ref_system_rows))
        vbox_ref_sys.addWidget(btn_add_ref_sys)
        layout_ref.addWidget(box_ref_sys)

        self.add_system_row(self.layout_ref_sys, self.ref_system_rows, default_name="System_1")

        box_ref_metrics = QGroupBox("Metrics & Models")
        grid_ref_metrics = QGridLayout(box_ref_metrics)

        # Standard metrics checked by default
        self.cb_bleu = QCheckBox("BLEU (SacreBLEU)")
        self.cb_bleu.setChecked(True)
        self.cb_chrf = QCheckBox("chrF++ (Character & Word n-grams)")
        self.cb_chrf.setChecked(True)
        self.cb_ter = QCheckBox("TER (Translation Edit Rate)")
        self.cb_ter.setChecked(True)
        self.cb_ed = QCheckBox("Edit Distance (%EdDist & WER)")
        self.cb_ed.setChecked(True)

        # Neural metrics UNCHECKED by default
        self.cb_comet = QCheckBox("COMET:")
        self.cb_comet.setChecked(False)
        self.txt_comet_model = QLineEdit(models_cfg.get("comet_default", ""))

        self.cb_bertscore = QCheckBox("BERTScore:")
        self.cb_bertscore.setChecked(False)
        self.txt_bertscore_model = QLineEdit(models_cfg.get("bertscore_default", ""))

        grid_ref_metrics.addWidget(self.cb_bleu, 0, 0)
        grid_ref_metrics.addWidget(self.cb_chrf, 0, 1)
        grid_ref_metrics.addWidget(self.cb_ter, 1, 0)
        grid_ref_metrics.addWidget(self.cb_ed, 1, 1)

        row_c = QHBoxLayout(); row_c.addWidget(self.cb_comet); row_c.addWidget(self.txt_comet_model)
        grid_ref_metrics.addLayout(row_c, 2, 0)

        row_b = QHBoxLayout(); row_b.addWidget(self.cb_bertscore); row_b.addWidget(self.txt_bertscore_model)
        grid_ref_metrics.addLayout(row_b, 2, 1)

        layout_ref.addWidget(box_ref_metrics)

        btn_run_ref = QPushButton("Run Reference-Based Evaluation")
        btn_run_ref.setStyleSheet("font-size: 13px; font-weight: bold; background-color: #0066cc; color: white; padding: 9px;")
        btn_run_ref.clicked.connect(lambda: self.start_evaluation("reference"))
        layout_ref.addWidget(btn_run_ref)

        self.tabs.addTab(tab_ref, "Reference-Based Evaluation")

        # -------------------------------------------------------------------
        # TAB 2: Referenceless Evaluation (Quality Estimation / QE)
        # -------------------------------------------------------------------
        tab_qe = QWidget()
        layout_qe = QVBoxLayout(tab_qe)

        box_qe_files = QGroupBox("Input File (No reference needed)")
        grid_qe_files = QGridLayout(box_qe_files)
        lbl_qe_src = QLabel("<b>Source file (* Required):</b>")
        grid_qe_files.addWidget(lbl_qe_src, 0, 0)
        self.txt_qe_src = QLineEdit()
        self.txt_qe_src.setPlaceholderText("Select source text file (*)")
        grid_qe_files.addWidget(self.txt_qe_src, 0, 1)
        btn_browse_qe_src = QPushButton("Browse...")
        btn_browse_qe_src.clicked.connect(lambda: self.browse_file(self.txt_qe_src, "Select Source File"))
        grid_qe_files.addWidget(btn_browse_qe_src, 0, 2)
        layout_qe.addWidget(box_qe_files)

        box_qe_sys = QGroupBox("MT Systems to Evaluate")
        vbox_qe_sys = QVBoxLayout(box_qe_sys)
        scroll_qe = QScrollArea()
        scroll_qe.setWidgetResizable(True)
        scroll_qe.setMinimumHeight(150)
        scroll_qe.setFrameShape(QFrame.Shape.NoFrame)
        self.cont_qe_sys = QWidget()
        self.layout_qe_sys = QVBoxLayout(self.cont_qe_sys)
        self.layout_qe_sys.setAlignment(Qt.AlignmentFlag.AlignTop)
        scroll_qe.setWidget(self.cont_qe_sys)
        vbox_qe_sys.addWidget(scroll_qe)

        btn_add_qe_sys = QPushButton("+ Add Another System")
        btn_add_qe_sys.setStyleSheet("font-weight: bold; padding: 6px;")
        btn_add_qe_sys.clicked.connect(lambda: self.add_system_row(self.layout_qe_sys, self.qe_system_rows))
        vbox_qe_sys.addWidget(btn_add_qe_sys)
        layout_qe.addWidget(box_qe_sys)

        self.add_system_row(self.layout_qe_sys, self.qe_system_rows, default_name="System_1")

        box_qe_metrics = QGroupBox("Quality Estimation Metrics")
        grid_qe_metrics = QGridLayout(box_qe_metrics)

        # COMET-QE checked by default; SBERT unchecked by default
        self.cb_comet_qe = QCheckBox("COMET-QE:")
        self.cb_comet_qe.setChecked(True)
        self.txt_comet_qe_model = QLineEdit(models_cfg.get("comet_qe_default", ""))
        row_qe = QHBoxLayout(); row_qe.addWidget(self.cb_comet_qe); row_qe.addWidget(self.txt_comet_qe_model)
        grid_qe_metrics.addLayout(row_qe, 0, 0)

        self.cb_sbert_qe = QCheckBox("SBERT (Cosine Sim):")
        self.cb_sbert_qe.setChecked(False)
        self.txt_sbert_qe_model = QLineEdit(models_cfg.get("sbert_qe_default", ""))
        row_sbert = QHBoxLayout(); row_sbert.addWidget(self.cb_sbert_qe); row_sbert.addWidget(self.txt_sbert_qe_model)
        grid_qe_metrics.addLayout(row_sbert, 1, 0)

        layout_qe.addWidget(box_qe_metrics)

        btn_run_qe = QPushButton("Run Referenceless (QE) Evaluation")
        btn_run_qe.setStyleSheet("font-size: 13px; font-weight: bold; background-color: #2e7d32; color: white; padding: 9px;")
        btn_run_qe.clicked.connect(lambda: self.start_evaluation("referenceless"))
        layout_qe.addWidget(btn_run_qe)

        self.tabs.addTab(tab_qe, "Referenceless (QE)")

        # -------------------------------------------------------------------
        # TAB 3: Segment-level Details (Ordered: ID, System, Metrics, Hyp, Ref, Src, Diff)
        # -------------------------------------------------------------------
        tab_details = QWidget()
        layout_details = QVBoxLayout(tab_details)
        self.table_segments = QTableWidget()
        self.table_segments.setAlternatingRowColors(True)
        layout_details.addWidget(self.table_segments)
        self.tabs.addTab(tab_details, "Segment Details")

        # -------------------------------------------------------------------
        # TAB 4: Visualizations (Individual Sub-Tabs per Metric and Radar)
        # -------------------------------------------------------------------
        self.vis_widget = VisualizationWidget()
        self.tabs.addTab(self.vis_widget, "Visualizations")

        # -------------------------------------------------------------------
        # TAB 5: Reproducibility Signatures
        # -------------------------------------------------------------------
        tab_sig = QWidget()
        layout_sig = QVBoxLayout(tab_sig)
        lbl_sig_desc = QLabel(
            "<b>Evaluation Signatures for Papers:</b><br>"
            "Copy and paste these exact signatures to guarantee reproducibility."
        )
        layout_sig.addWidget(lbl_sig_desc)

        self.txt_signatures = QPlainTextEdit()
        self.txt_signatures.setReadOnly(True)
        self.txt_signatures.setStyleSheet("font-family: Consolas, monospace; font-size: 11px; background-color: #f7f7f7;")
        layout_sig.addWidget(self.txt_signatures)

        btn_copy_sig = QPushButton("Copy Signatures to Clipboard")
        btn_copy_sig.clicked.connect(self.copy_signatures)
        layout_sig.addWidget(btn_copy_sig)

        self.tabs.addTab(tab_sig, "Signatures")

        # -------------------------------------------------------------------
        # Bottom Global Area
        # -------------------------------------------------------------------
        self.status_lbl = QLabel("Ready.")
        main_layout.addWidget(self.status_lbl)

        self.progress_bar = QProgressBar()
        self.progress_bar.setValue(0)
        main_layout.addWidget(self.progress_bar)

        h_exports = QHBoxLayout()
        self.btn_export_latex = QPushButton("Export to LaTeX...")
        self.btn_export_latex.setEnabled(False)
        self.btn_export_latex.clicked.connect(self.export_latex)
        h_exports.addWidget(self.btn_export_latex)

        self.btn_export_excel = QPushButton("Export to Excel (.xlsx)")
        self.btn_export_excel.setEnabled(False)
        self.btn_export_excel.clicked.connect(self.export_excel)
        h_exports.addWidget(self.btn_export_excel)

        self.btn_export_csv = QPushButton("Export to CSV (.csv)")
        self.btn_export_csv.setEnabled(False)
        self.btn_export_csv.clicked.connect(self.export_csv)
        h_exports.addWidget(self.btn_export_csv)
        main_layout.addLayout(h_exports)

        self.table_corpus = QTableWidget()
        main_layout.addWidget(self.table_corpus)

    def add_system_row(self, layout: QVBoxLayout, storage: List[SystemRowWidget], default_name: str = ""):
        idx = len(storage) + 1
        widget = SystemRowWidget(idx, default_name=default_name)
        widget.remove_requested.connect(lambda w: self.remove_system_row(w, layout, storage))
        layout.addWidget(widget)
        storage.append(widget)

    def remove_system_row(self, widget: SystemRowWidget, layout: QVBoxLayout, storage: List[SystemRowWidget]):
        if len(storage) <= 1:
            QMessageBox.warning(self, "Warning", "At least one system is required.")
            return
        layout.removeWidget(widget)
        storage.remove(widget)
        widget.deleteLater()
        for i, row in enumerate(storage):
            row.lbl.setText(f"System #{i + 1}:")

    def browse_file(self, target_line: QLineEdit, title: str):
        path, _ = QFileDialog.getOpenFileName(self, title, "", "All Files (*)")
        if path:
            target_line.setText(path)

    def start_evaluation(self, mode: str):
        if mode == "reference":
            ref_path = self.txt_ref_ref.text().strip()
            src_path = self.txt_ref_src.text().strip()
            rows = self.ref_system_rows
            selected_metrics = {
                "bleu": self.cb_bleu.isChecked(),
                "chrf": self.cb_chrf.isChecked(),
                "ter": self.cb_ter.isChecked(),
                "ed_dist": self.cb_ed.isChecked(),
                "comet": self.cb_comet.isChecked(),
                "bertscore": self.cb_bertscore.isChecked(),
            }
        else:
            ref_path = ""
            src_path = self.txt_qe_src.text().strip()
            rows = self.qe_system_rows
            selected_metrics = {
                "comet_qe": self.cb_comet_qe.isChecked(),
                "sbert_qe": self.cb_sbert_qe.isChecked(),
            }

        systems_data = []
        for r in rows:
            d = r.get_data()
            if not d["path"]:
                continue
            if not os.path.exists(d["path"]):
                QMessageBox.critical(self, "File Not Found", f"File not found for system '{d['name']}':\n{d['path']}")
                return
            systems_data.append(d)

        if not systems_data:
            QMessageBox.critical(self, "Error", "Please configure at least one system with a valid file.")
            return

        if not any(selected_metrics.values()):
            QMessageBox.warning(self, "No Metrics", "Please select at least one evaluation metric.")
            return

        self.btn_export_latex.setEnabled(False)
        self.btn_export_excel.setEnabled(False)
        self.btn_export_csv.setEnabled(False)
        self.table_corpus.clear()
        self.table_segments.clear()
        self.txt_signatures.clear()

        models_cfg = self.config.get("Models", {})
        c_model = self.txt_comet_model.text().strip() or models_cfg.get("comet_default", "")
        cqe_model = self.txt_comet_qe_model.text().strip() or models_cfg.get("comet_qe_default", "")
        bs_model = self.txt_bertscore_model.text().strip() or models_cfg.get("bertscore_default", "")
        sb_model = self.txt_sbert_qe_model.text().strip() or models_cfg.get("sbert_qe_default", "")

        self.worker = EvaluationWorker(
            mode=mode,
            src_path=src_path,
            ref_path=ref_path,
            systems=systems_data,
            selected_metrics=selected_metrics,
            config=self.config,
            comet_model=c_model,
            comet_qe_model=cqe_model,
            bertscore_model=bs_model,
            sbert_qe_model=sb_model
        )
        self.worker.progress_signal.connect(self.on_progress)
        self.worker.finished_signal.connect(self.on_finished)
        self.worker.error_signal.connect(self.on_error)
        self.worker.start()

    def on_progress(self, msg: str, val: int):
        self.status_lbl.setText(msg)
        self.progress_bar.setValue(val)

    def on_finished(self, results: dict):
        self.results_data = results
        self.status_lbl.setText("Evaluation completed.")
        self.progress_bar.setValue(100)
        self.btn_export_latex.setEnabled(True)
        self.btn_export_excel.setEnabled(True)
        self.btn_export_csv.setEnabled(True)

        corpus = results["corpus"]
        systems = results["systems"]
        metrics = list(corpus.keys())

        # 1. Corpus Summary Table
        self.table_corpus.setColumnCount(len(systems) + 1)
        self.table_corpus.setHorizontalHeaderLabels(["Metric"] + systems)
        self.table_corpus.setRowCount(len(metrics))

        for r_idx, m_name in enumerate(metrics):
            self.table_corpus.setItem(r_idx, 0, QTableWidgetItem(m_name))
            for c_idx, s_name in enumerate(systems):
                score = corpus[m_name].get(s_name, "")
                item = QTableWidgetItem(str(score))
                item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                self.table_corpus.setItem(r_idx, c_idx + 1, item)

        self.table_corpus.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)

        # 2. Detailed Segment Table: ID, SYSTEM, [METRICS], HYPOTHESIS, REFERENCE, ORIGINAL, DIFF
        segments = results["segments"]
        if segments:
            sample_metrics = list(segments[0]["metrics"].keys())
            headers = ["ID", "SYSTEM"] + [m.upper() for m in sample_metrics] + ["HYPOTHESIS"]
            if results["has_ref"]:
                headers.append("REFERENCE")
            if results["has_source"]:
                headers.append("ORIGINAL")
            if results["has_ref"]:
                headers.append("DIFF")

            self.table_segments.setColumnCount(len(headers))
            self.table_segments.setHorizontalHeaderLabels(headers)
            self.table_segments.setRowCount(len(segments))

            for r_idx, rec in enumerate(segments):
                c_idx = 0
                self.table_segments.setItem(r_idx, c_idx, QTableWidgetItem(str(rec["seg_id"]))); c_idx += 1
                self.table_segments.setItem(r_idx, c_idx, QTableWidgetItem(rec["system"])); c_idx += 1

                for m in sample_metrics:
                    val = rec["metrics"].get(m, "")
                    item = QTableWidgetItem(str(val))
                    item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                    self.table_segments.setItem(r_idx, c_idx, item)
                    c_idx += 1

                self.table_segments.setItem(r_idx, c_idx, QTableWidgetItem(rec["hypothesis"])); c_idx += 1
                if results["has_ref"]:
                    self.table_segments.setItem(r_idx, c_idx, QTableWidgetItem(rec["reference"])); c_idx += 1
                if results["has_source"]:
                    self.table_segments.setItem(r_idx, c_idx, QTableWidgetItem(rec["source"])); c_idx += 1
                if results["has_ref"]:
                    self.table_segments.setItem(r_idx, c_idx, QTableWidgetItem(rec["diff"]))

            self.table_segments.resizeColumnsToContents()

        # 3. Update Visualizations (Tabs for Bar charts + Tabs for Radar charts)
        self.vis_widget.update_plots(corpus, systems)

        # 4. Signatures Tab
        signatures = results.get("signatures", {})
        sig_text_lines = ["# Evaluation Signatures & Reproducibility Metadata\n"]
        for m, sig in signatures.items():
            sig_text_lines.append(f"{m}:\n  {sig}\n")
        self.txt_signatures.setPlainText("\n".join(sig_text_lines))

        QMessageBox.information(self, "Done", "Evaluation finished successfully!")

    def on_error(self, err_msg: str):
        self.status_lbl.setText("Error.")
        self.progress_bar.setValue(0)
        QMessageBox.critical(self, "Evaluation Error", f"An error occurred:\n{err_msg}")

    def copy_signatures(self):
        text = self.txt_signatures.toPlainText()
        if text:
            QApplication.clipboard().setText(text)
            QMessageBox.information(self, "Copied", "Signatures copied to clipboard!")

    # -----------------------------------------------------------------------
    # Export Options: LaTeX, Excel & CSV
    # -----------------------------------------------------------------------

    def export_latex(self):
        if not self.results_data or not self.results_data.get("corpus"):
            return

        corpus = self.results_data["corpus"]
        systems = self.results_data["systems"]
        metrics = list(corpus.keys())

        col_def = "l" + "c" * len(systems)
        latex_lines = [
            "% Add \\usepackage{booktabs} to your LaTeX preamble",
            "\\begin{table}[htbp]",
            "\\centering",
            "\\caption{Machine Translation evaluation results across systems.}",
            "\\label{tab:mt_evaluation_results}",
            f"\\begin{{tabular}}{{{col_def}}}",
            "\\toprule",
            "\\textbf{Metric} & " + " & ".join(f"\\textbf{{{s}}}" for s in systems) + " \\\\",
            "\\midrule"
        ]

        for m in metrics:
            row_vals = []
            for s in systems:
                val = corpus[m].get(s, "")
                row_vals.append(str(val))
            escaped_metric = m.replace("%", "\\%")
            latex_lines.append(f"{escaped_metric} & " + " & ".join(row_vals) + " \\\\")

        latex_lines.extend([
            "\\bottomrule",
            "\\end{tabular}",
            "\\end{table}"
        ])

        dlg = LatexExportDialog("\n".join(latex_lines), parent=self)
        dlg.exec()

    def export_excel(self):
        if not self.results_data or not HAS_OPENPYXL:
            if not HAS_OPENPYXL:
                QMessageBox.warning(self, "Missing Library", "Install openpyxl via: pip install openpyxl")
            return

        path, _ = QFileDialog.getSaveFileName(self, "Export Results to Excel", "mt_evaluation_report.xlsx", "All Files (*)")
        if not path:
            return
        if not path.endswith(".xlsx"):
            path += ".xlsx"

        try:
            wb = openpyxl.Workbook()
            header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

            ws_summary = wb.active
            ws_summary.title = "Corpus Summary"
            ws_summary.views.sheetView[0].showGridLines = True

            systems = self.results_data["systems"]
            corpus = self.results_data["corpus"]

            summary_headers = ["Metric"] + systems
            ws_summary.append(summary_headers)

            for col_idx in range(1, len(summary_headers) + 1):
                cell = ws_summary.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left")

            for metric_name, sys_scores in corpus.items():
                row = [metric_name] + [sys_scores.get(s, "") for s in systems]
                ws_summary.append(row)

            ws_det = wb.create_sheet(title="Segment Details")
            ws_det.views.sheetView[0].showGridLines = True

            segments = self.results_data["segments"]
            sample_metrics = list(segments[0]["metrics"].keys()) if segments else []
            has_source = self.results_data["has_source"]
            has_ref = self.results_data["has_ref"]

            # Order: ID, System, [Metrics], Hypothesis, Reference, Original, Diff
            det_headers = ["ID", "System"] + sample_metrics + ["Hypothesis"]
            if has_ref:
                det_headers.append("Reference")
            if has_source:
                det_headers.append("Original")
            if has_ref:
                det_headers.append("DIFF")

            ws_det.append(det_headers)

            for col_idx in range(1, len(det_headers) + 1):
                cell = ws_det.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font

            for rec in segments:
                row_vals = [rec["seg_id"], rec["system"]]
                for m in sample_metrics:
                    row_vals.append(rec["metrics"].get(m, ""))
                row_vals.append(rec["hypothesis"])
                if has_ref:
                    row_vals.append(rec["reference"])
                if has_source:
                    row_vals.append(rec["source"])
                if has_ref:
                    row_vals.append(rec["diff"])
                ws_det.append(row_vals)

            wb.save(path)
            QMessageBox.information(self, "Exported", f"Excel report saved to:\n{path}")

        except Exception as ex:
            QMessageBox.critical(self, "Export Error", f"Failed to export Excel:\n{ex}")

    def export_csv(self):
        if not self.results_data or not self.results_data.get("segments"):
            return

        path, _ = QFileDialog.getSaveFileName(self, "Export Segments to CSV", "mt_evaluation_segments.csv", "All Files (*)")
        if not path:
            return
        if not path.endswith(".csv"):
            path += ".csv"

        try:
            segments = self.results_data["segments"]
            sample_metrics = list(segments[0]["metrics"].keys())
            has_source = self.results_data["has_source"]
            has_ref = self.results_data["has_ref"]

            # Order: ID, System, [Metrics], Hypothesis, Reference, Original, Diff
            headers = ["ID", "System"] + sample_metrics + ["Hypothesis"]
            if has_ref:
                headers.append("Reference")
            if has_source:
                headers.append("Original")
            if has_ref:
                headers.append("DIFF")

            with open(path, "w", encoding="utf-8-sig", newline="") as f:
                writer = csv.writer(f)
                writer.writerow(headers)

                for rec in segments:
                    row = [rec["seg_id"], rec["system"]]
                    for m in sample_metrics:
                        row.append(rec["metrics"].get(m, ""))
                    row.append(rec["hypothesis"])
                    if has_ref:
                        row.append(rec["reference"])
                    if has_source:
                        row.append(rec["source"])
                    if has_ref:
                        row.append(rec["diff"])
                    writer.writerow(row)

            QMessageBox.information(self, "Exported", f"CSV file saved to:\n{path}")

        except Exception as ex:
            QMessageBox.critical(self, "Export Error", f"Failed to export CSV:\n{ex}")


# ---------------------------------------------------------------------------
# Application Entry Point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = MTUOCEvalApp()
    window.show()
    sys.exit(app.exec())
